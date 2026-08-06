from pathlib import Path
from docx import Document

path = Path(r"G:\My Drive\Coding in r (lakwong@stanford.edu)\Rohingya_analysis\7_tables\RF105_reviewed_20260717\RF105_original_vs_reviewed_results_summary.docx")
doc = Document(path)

print(f"PARAGRAPHS: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs[:80], start=1):
    text = p.text.strip()
    if text:
        print(f"P{i:03d} [{p.style.name}]: {text}")

print(f"\nTABLES: {len(doc.tables)}")
for ti, table in enumerate(doc.tables, start=1):
    print(f"\nTABLE {ti}: {len(table.rows)} rows x {len(table.columns)} cols")
    for ri, row in enumerate(table.rows[:8], start=1):
        vals = [cell.text.replace("\n", " | ").strip() for cell in row.cells]
        print(f"R{ri:02d}: " + " || ".join(vals))
